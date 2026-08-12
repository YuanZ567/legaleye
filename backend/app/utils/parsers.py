"""文档解析工具：文本提取（PDF/Word/HTML）+ 扩展名识别 + docType 启发式推断。

对应 PRD F3：
- PyMuPDF 提取 PDF 文本；扫描版（无文字层）抛出专用异常（E2）；
- python-docx 提取 Word（.docx）文本；
- BeautifulSoup 提取 HTML 文本（URL 解析）。
"""

from pathlib import Path

from app.core.enums import DocType


class ScannedPdfError(Exception):
    """扫描版 PDF（无文字层）专用异常，映射 E2 语义。"""


def get_file_extension(filename: str) -> str:
    """返回小写扩展名（含点），如 'foo.PDF' → '.pdf'。"""
    return Path(filename).suffix.lower()


def is_pdf(filename: str) -> bool:
    return get_file_extension(filename) == ".pdf"


def count_pdf_pages(file_bytes: bytes) -> int:
    """读取 PDF 页数（用于 ≤200 页校验）。非 PDF 文件调用方不得调用。"""
    import pymupdf  # PyMuPDF（延迟导入，避免未使用路径拉高启动开销）

    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        return doc.page_count


def extract_pdf_text(file_bytes: bytes) -> str:
    """提取 PDF 全文（PRD F3）。

    若所有页均无文字层（扫描版），抛 ScannedPdfError（E2：MVP 不支持 OCR）。
    """
    import pymupdf

    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        pages_text = [page.get_text().strip() for page in doc]
    combined = "\n".join(pages_text).strip()
    if not combined:
        raise ScannedPdfError("未检测到文字层，MVP 不支持 OCR，请提供文字版 PDF")
    return combined


def extract_docx_text(file_bytes: bytes) -> str:
    """提取 Word（.docx）全文（PRD F3）。"""
    import io

    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs]
    # 表格文本也纳入（合规文件常含数据清单）
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(p for p in paragraphs if p)


# docType 文件名启发式推断（PRD F3：未选择时按文件名推断，失败提示手动选择）
_DOC_TYPE_FILENAME_HINTS: dict[DocType, tuple[str, ...]] = {
    DocType.PRIVACY_POLICY: ("隐私政策", "隐私", "privacy", "policy"),
    DocType.USER_AGREEMENT: ("用户协议", "用户条款", "agreement", "terms", "user"),
    DocType.DPA: ("数据处理协议", "处理协议", "dpa", "data process"),
    DocType.SCC: ("标准合同", "标准条款", "scc", "standard contractual"),
}


def infer_doc_type_from_filename(filename: str) -> DocType | None:
    """按文件名关键词推断 docType；无法确定返回 None（提示手动选择）。"""
    lower_name = filename.lower()
    for doc_type, hints in _DOC_TYPE_FILENAME_HINTS.items():
        if any(hint in lower_name for hint in hints):
            return doc_type
    return None


def html_to_text(html: str) -> str:
    """将 HTML 转为纯文本（URL 解析用，PRD F3：抓取 html 转文本）。

    保留脚本/样式外的正文文字，多空白归一化为单空格。
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    # 移除脚本与样式，避免混入非正文内容
    for node in soup(["script", "style", "noscript"]):
        node.decompose()
    text = soup.get_text(separator=" ")
    return " ".join(text.split())
